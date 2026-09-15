from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path('docs/phan-tich-ket-qua-danh-gia-atad-de-hieu.docx')


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    cant_split.set(qn('w:val'), 'true')
    tr_pr.append(cant_split)


def set_col_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def style_table(table, header=True, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    if widths:
        set_col_widths(table, widths)
    for i, row in enumerate(table.rows):
        set_row_cant_split(row)
        if header and i == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = 'Aptos'
                    run.font.size = Pt(9.5)
            if header and i == 0:
                set_cell_shading(cell, '1F4E79')
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif i % 2 == 0:
                set_cell_shading(cell, 'F4F7FA')


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
    style_table(table, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    return p


def add_manual_number(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(f'{number}. ')
    r.bold = True
    p.add_run(text)
    return p


def add_note(doc, lead, text):
    p = doc.add_paragraph()
    r = p.add_run(lead)
    r.bold = True
    p.add_run(text)
    return p


def set_run_font(run, name='Aptos', size=10.5, color=None, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in [('Title', 23, (31, 78, 121)), ('Heading 1', 16, (31, 78, 121)), ('Heading 2', 12.5, (47, 84, 150)), ('Heading 3', 11.2, (68, 68, 68))]:
        st = styles[style_name]
        st.font.name = 'Aptos Display' if style_name in ('Title', 'Heading 1') else 'Aptos'
        st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
        st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(*color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(10 if style_name != 'Title' else 0)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True

    # Cover and opening conclusion
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run('Phân tích kết quả đánh giá website ATAD')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run('Diễn giải dễ hiểu từ 5 lần đánh giá bằng WebLens')
    set_run_font(r, size=12, color=(90, 90, 90))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Kết luận ngắn gọn: ')
    r.bold = True
    p.add_run('Website ATAD có một số vấn đề đáng chú ý về bố cục, trải nghiệm trên màn hình nhỏ, khả năng tiếp cận và lỗi khi chạy trang. Các điểm số của bộ đánh giá khá ổn định và bằng chứng được trích dẫn đầy đủ. Tuy vậy, bộ đánh giá chưa đạt ngưỡng tin cậy tổng thể vì các phát hiện do mô hình tự mô tả lại không được nhận diện giống nhau qua các lần chạy. Vì thế, kết quả đủ tốt để xác định các hướng cần kiểm tra và cải thiện, nhưng chưa nên xem là kết luận tuyệt đối hoặc thay thế cho việc kiểm tra của con người.')

    add_note(doc, 'Đối tượng của tài liệu: ', 'người cần hiểu kết quả đánh giá website nhưng không cần biết lập trình hay các chi tiết kỹ thuật của hệ thống.')
    add_note(doc, 'Dữ liệu sử dụng: ', '5 lần chạy ATAD đã hoàn tất và được chấm lại ở chế độ offline. Mỗi lần gồm thu thập trang, chạy kiểm tra tự động và để mô hình chấm theo cùng một bộ tiêu chí.')
    add_note(doc, 'Phạm vi: ', 'đánh giá trang web và độ ổn định của bộ máy đánh giá; không phải kiểm định chất lượng kinh doanh, nội dung pháp lý hay mức độ đúng của mọi thông tin trên website.')

    doc.add_heading('1. Cách đọc tài liệu này', level=1)
    doc.add_paragraph('Có hai câu hỏi khác nhau cần tách riêng:')
    add_number(doc, 'Website ATAD đang có điểm mạnh, điểm yếu nào? Đây là phần “kết quả đánh giá website”.')
    add_number(doc, 'Có thể tin bộ máy đánh giá đến mức nào? Đây là phần “độ tin cậy của bộ đánh giá”.')
    doc.add_paragraph('Một bộ máy có thể chấm rất đều nhưng vẫn chấm sai; ngược lại, một phát hiện có thể đúng nhưng cách mô hình đặt tên lại thay đổi giữa các lần chạy. Báo cáo này phân biệt hai việc đó để tránh hiểu nhầm.')

    doc.add_heading('2. Tóm tắt kết quả ATAD trong một trang', level=1)
    add_table(doc, ['Nội dung', 'Kết quả', 'Diễn giải đời thường'], [
        ('Số lần đánh giá', '5/5 lần hoàn tất', 'Đủ số mẫu tối thiểu để xem xét độ ổn định.'),
        ('Điểm thị giác', '35/100', 'Bố cục, khoảng cách và thứ bậc nhìn còn yếu theo rubric của hệ thống.'),
        ('Điểm trải nghiệm', '53,2/100', 'Cấu trúc thông tin, nút hành động và khả năng đọc ở mức trung bình.'),
        ('Điểm responsive', '45/100', 'Có dấu hiệu hiển thị chưa tốt trên màn hình nhỏ; kết quả dao động giữa các lần.'),
        ('Điểm chuyển đổi', '50/100', 'Có thông tin liên hệ nhưng hành động chính chưa đủ nổi bật.'),
        ('Điểm kỹ thuật', '94,2/100', 'Kết quả Lighthouse trong môi trường kiểm thử khá cao; không có nghĩa website tổng thể đã tốt.'),
        ('Bằng chứng', '100% tồn tại và được hỗ trợ', 'Không thấy dấu hiệu trích dẫn tệp bằng chứng không có thật trong dữ liệu đã chấm.'),
        ('Phủ đủ tiêu chí', '100%', 'Đủ 13 tiêu chí rubric và đúng cấu trúc yêu cầu.'),
        ('Finding tái diễn', '44,5%', 'Nhiều phát hiện không được gọi tên giống nhau giữa 5 lần chạy.'),
        ('Trust gate', 'Không đạt', 'Bộ đánh giá chưa đủ ổn định để tự động dùng làm cổng quyết định.'),
    ], widths=[1.4, 1.35, 4.65])

    add_note(doc, 'Thông điệp quan trọng nhất: ', '“Trust gate không đạt” không đồng nghĩa website ATAD không có vấn đề. Nó chủ yếu cho thấy cách hệ thống gom và đặt tên các phát hiện bằng ngôn ngữ tự do chưa đủ ổn định.')

    doc.add_heading('3. Website ATAD đang được máy đánh giá như thế nào', level=1)
    doc.add_paragraph('Các điểm dưới đây là điểm của website theo rubric trong repo. Thang 0–100 dùng để dễ đọc; đây không phải điểm chuẩn ngành và không phải ý kiến của người dùng thật.')
    add_table(doc, ['Chiều đánh giá', 'Điểm trung bình', 'Khoảng giữa 5 lần', 'Diễn giải dễ hiểu'], [
        ('Thị giác', '35,0', '31–38', 'Thứ bậc nhìn và độ hoàn thiện còn yếu; menu và các khoảng trống lớn làm loãng màn hình đầu tiên.'),
        ('Trải nghiệm người dùng', '53,2', '50–58', 'Cấu trúc thông tin, nút hành động và khả năng đọc ở mức trung bình.'),
        ('Responsive', '45,0', '25–50', 'Có lần bị chấm thấp hơn rõ rệt, cho thấy cần xem kỹ giao diện mobile và khoảng trống đầu trang.'),
        ('Chuyển đổi', '50,0', '50–50', 'Điểm ổn định nhưng chỉ ở mức trung bình; lời mời liên hệ hoặc tư vấn chưa nổi bật.'),
        ('Kỹ thuật', '94,2', '93–96', 'Điểm kiểm thử hiệu năng trong phòng thí nghiệm cao; không bao phủ toàn bộ chất lượng trải nghiệm.'),
        ('Thương hiệu', 'Chưa chấm', '—', 'Repo chưa có rubric riêng để quy đổi thành điểm.'),
        ('Khả năng tiếp cận', 'Chưa quy đổi thành điểm', '—', 'Các lỗi vẫn được liệt kê bằng audit tự động, nhưng chưa gộp thành một điểm tổng.'),
    ], widths=[1.35, 1.05, 1.15, 3.85])

    doc.add_heading('3.1. Những điểm yếu nổi bật của website', level=2)
    add_bullet(doc, 'Trên desktop, menu lớn hoặc các mảng điều hướng màu xanh có thể chiếm phần lớn vùng nhìn đầu tiên, làm nội dung chính và hành động quan trọng bị lùi xuống.')
    add_bullet(doc, 'Trên mobile, có dấu hiệu khoảng trống dọc lớn trước phần nội dung; điều này làm người xem phải cuộn nhiều hơn mới thấy thông tin có giá trị.')
    add_bullet(doc, 'Ngôn ngữ hình ảnh giữa desktop và mobile chưa thật sự nhất quán; đây là lý do điểm visual.consistency chỉ đạt trung bình thấp.')
    add_bullet(doc, 'Thông tin liên hệ có xuất hiện, nhưng hành động chính như yêu cầu tư vấn hoặc bắt đầu trao đổi chưa đủ nổi bật để dẫn người xem đi tiếp.')
    add_bullet(doc, 'Khả năng đọc trên mobile là tiêu chí dao động nhiều nhất: điểm 2 xuất hiện 3 lần, điểm 3 xuất hiện 2 lần. Điều này cho thấy nhận xét là hợp lý để kiểm tra thêm, nhưng chưa nên xem là đo lường tuyệt đối.')

    doc.add_heading('3.2. Các phát hiện kỹ thuật lặp lại nhiều lần', level=2)
    doc.add_paragraph('Đây là nhóm kết quả có giá trị hành động cao hơn vì được công cụ tự động hoặc dữ liệu runtime ghi nhận lặp lại, ít phụ thuộc vào cách mô hình diễn đạt:')
    add_table(doc, ['Phát hiện', 'Số lần lặp lại', 'Ý nghĩa cần kiểm tra'], [
        ('heading-order', '5/5', 'Thứ tự các tiêu đề trên trang có thể không hợp lý đối với trình đọc màn hình.'),
        ('image-alt', '5/5', 'Một số hình ảnh có thể thiếu văn bản thay thế.'),
        ('landmark-one-main', '5/5', 'Cấu trúc vùng nội dung chính chưa rõ ràng đối với công cụ hỗ trợ.'),
        ('link-in-text-block', '5/5', 'Một số liên kết trong đoạn văn có thể chưa đáp ứng quy tắc kiểm tra.'),
        ('link-name', '5/5', 'Có liên kết chưa có tên đủ rõ cho người dùng hoặc trình đọc màn hình.'),
        ('page-has-heading-one', '5/5', 'Trang có thể thiếu hoặc dùng chưa đúng tiêu đề cấp một.'),
        ('region', '5/5', 'Một số vùng nội dung chưa được gắn cấu trúc rõ ràng.'),
        ('color-contrast', '5/5 trên mobile', 'Có khu vực cần kiểm tra lại độ tương phản màu.'),
        ('page-error', '5/5 desktop và 5/5 mobile', 'Trang phát sinh lỗi runtime trong lúc thu thập; cần xem log trình duyệt và mã JavaScript.'),
        ('horizontal-overflow', '4/5 desktop', 'Có lần nội dung rộng hơn khung nhìn, gây tràn ngang.'),
    ], widths=[1.65, 1.2, 4.55])
    add_note(doc, 'Cách ưu tiên: ', 'các lỗi lặp lại 5/5 nên được kiểm tra trước. Riêng page-error cần điều tra nguyên nhân cụ thể, vì điểm Lighthouse cao không loại trừ lỗi JavaScript khi người dùng thật truy cập.')

    doc.add_heading('4. Mô hình đã chấm từng tiêu chí ra sao', level=1)
    doc.add_paragraph('Mỗi tiêu chí được chấm từ 0 đến 4. Điểm càng cao nghĩa là tiêu chí càng đạt yêu cầu theo rubric. Đây là thang đánh giá nội bộ của repo, không phải thang điểm phổ quát.')
    add_table(doc, ['Nhóm', 'Tiêu chí', 'Điểm trung bình', 'Mức ổn định', 'Diễn giải'], [
        ('Thị giác', 'visual.hierarchy', '1,0/4', '100% giống mẫu đầu', 'Thứ bậc nhìn bị menu hoặc vùng điều hướng lấn át.'),
        ('Thị giác', 'visual.composition', '1,0/4', '100%', 'Bố cục và nhịp nội dung chưa gọn.'),
        ('Thị giác', 'visual.typography', '2,0/4', '100%', 'Kiểu chữ ở mức trung bình.'),
        ('Thị giác', 'visual.spacing', '1,0/4', '100%', 'Khoảng cách và khoảng trống chưa được cân bằng tốt.'),
        ('Thị giác', 'visual.color', '2,0/4', '100%', 'Màu sắc ở mức trung bình.'),
        ('Thị giác', 'visual.assets', '1,8/4', '80%', 'Hình ảnh, biểu tượng hoặc tài sản trực quan dao động nhẹ.'),
        ('Thị giác', 'visual.consistency', '1,4/4', '60%', 'Desktop và mobile chưa thống nhất về ngôn ngữ hình ảnh.'),
        ('Thị giác', 'visual.polish', '1,0/4', '100%', 'Chi tiết hoàn thiện cuối cùng còn yếu.'),
        ('Trải nghiệm', 'ux.information-architecture', '2,0/4', '100%', 'Cấu trúc thông tin ở mức trung bình.'),
        ('Trải nghiệm', 'ux.cta-clarity', '2,0/4', '100%', 'Nút hoặc hành động chính có thể hiểu được nhưng chưa nổi bật.'),
        ('Trải nghiệm', 'ux.readability', '2,4/4', '40%', 'Khả năng đọc là tiêu chí dao động nhất.'),
        ('Responsive', 'responsive.behavior', '1,8/4', '80%', 'Phần lớn lần chấm ở mức 2, có một lần ở mức 1.'),
        ('Chuyển đổi', 'conversion.affordances', '2,0/4', '100%', 'Có đường dẫn liên hệ nhưng chưa dẫn dắt mạnh.'),
    ], widths=[1.0, 1.8, 0.9, 1.0, 2.7])
    add_note(doc, 'Cần hiểu đúng: ', 'điểm thấp ở visual hoặc UX không tự động chứng minh website “tệ”. Nó chỉ nói rằng theo các tiêu chí được định nghĩa trong repo, những khu vực đó đang có tín hiệu cần xem xét.')

    doc.add_heading('5. Bộ máy đánh giá có đáng tin không', level=1)
    doc.add_paragraph('Phần này không chấm website. Phần này kiểm tra xem kết quả do bộ máy tạo ra có đủ ổn định, đủ bằng chứng và đúng cấu trúc hay không.')
    add_table(doc, ['Kiểm tra', 'Kết quả', 'Đánh giá dễ hiểu'], [
        ('Đủ số lần chạy', '5/5', 'Đạt số mẫu tối thiểu.'),
        ('Lần chạy hoàn tất và chấm được', '5/5', 'Không có lần nào bị lỗi kỹ thuật hoặc không thể chấm.'),
        ('Độ dao động điểm tiêu chí', 'SD trung bình 0,14', 'Tốt; thấp hơn ngưỡng 0,8.'),
        ('Đồng thuận đúng tuyệt đối', '89,2%', 'Phần lớn tiêu chí lặp lại đúng cùng một điểm.'),
        ('Đồng thuận trong lệch một bậc', '100%', 'Không tiêu chí nào lệch quá một bậc.'),
        ('Tồn tại bằng chứng', '100%', 'Các mã bằng chứng được dẫn đều tồn tại trong dữ liệu.'),
        ('Bằng chứng hỗ trợ finding', '100%', 'Các finding có thể đánh giá đều có dữ liệu hỗ trợ.'),
        ('Tỷ lệ claim không được hỗ trợ', '0%', 'Không thấy claim bị gắn là không có bằng chứng trong mẫu này.'),
        ('Đủ rubric và schema', '100%', 'Đủ 13 tiêu chí, đúng cấu trúc, đủ trường bắt buộc.'),
        ('Tính hợp lệ số liệu audit', '100%', 'Các số Lighthouse và audit nằm trong miền hợp lệ.'),
        ('Finding tái diễn', '44,5%', 'Không đạt ngưỡng 80%; đây là điểm làm trust gate thất bại.'),
        ('Đối chiếu với người', 'Chưa có', 'Chưa thể kết luận độ chính xác tuyệt đối.'),
    ], widths=[1.85, 1.25, 4.3])

    doc.add_heading('5.1. Vì sao trust gate không đạt', level=2)
    doc.add_paragraph('Trong 5 lần chạy, hệ thống tạo ra 58 nhóm finding khác nhau. Chỉ 18 nhóm xuất hiện đủ thường xuyên để được xem là ổn định; 40 nhóm còn lại xuất hiện ít hơn ngưỡng yêu cầu. Tỷ lệ tái diễn trung bình chỉ là 44,5%, trong khi ngưỡng cần đạt là 80%.')
    doc.add_paragraph('Nguyên nhân chính không nhất thiết là mô hình nhìn sai. Cùng một hiện tượng, chẳng hạn menu chiếm quá nhiều vùng nhìn đầu tiên, có thể bị mô hình đặt các tên khác nhau như “navigation obstruction”, “mega menu overload” hoặc “navigation dominates”. Hệ thống hiện đang ghép finding bằng mã và tên gần như chính xác tuyệt đối. Vì vậy, cùng một ý nhưng khác cách viết bị tính thành nhiều finding khác nhau.')
    add_note(doc, 'Kết luận đúng: ', 'bộ máy chưa ổn định ở cách nhận diện và đặt tên finding. Không được diễn giải thành “AI bịa ra 55,5% vấn đề”.')

    doc.add_heading('5.2. Hallucination được kiểm tra đến đâu', level=2)
    doc.add_paragraph('Trong tài liệu này, hallucination được đo bằng một chỉ báo gián tiếp: finding có trỏ tới bằng chứng tồn tại và được dữ liệu hỗ trợ hay không. Kết quả là tồn tại bằng chứng 100%, hỗ trợ bằng chứng 100% và claim không được hỗ trợ 0%. Đây là tín hiệu tốt.')
    doc.add_paragraph('Tuy nhiên, các tỷ lệ này chưa phải là phép chứng minh mọi câu mô tả của mô hình đều đúng với thực tế. Một finding chủ quan có thể có ảnh đi kèm nhưng cách diễn giải vẫn cần người kiểm tra. Muốn đo “đúng như người”, repo còn thiếu bộ nhãn chuẩn do ít nhất hai người đánh giá độc lập tạo ra.')

    doc.add_heading('5.3. Classification ổn định nhưng design language chưa ổn định', level=2)
    add_table(doc, ['Nội dung', 'Kết quả', 'Ý nghĩa'], [
        ('Archetype', 'corporate_service, đồng thuận 100%', 'Mô hình nhận diện ATAD là website dịch vụ doanh nghiệp khá nhất quán.'),
        ('Design language', 'Đồng thuận chuỗi chính xác 6,7%', 'Cách viết mô tả phong cách thiết kế thay đổi nhiều; không nên dùng exact-match để kết luận mô hình không hiểu website.'),
        ('Mức tự tin của mô hình', '15/15 ở mức high', 'Mô hình tự tin cao, nhưng tự tin không đồng nghĩa đúng tuyệt đối.'),
    ], widths=[1.65, 1.6, 4.15])

    doc.add_page_break()
    doc.add_heading('6. Điều gì có thể dùng ngay và điều gì chưa nên dùng', level=1)
    add_table(doc, ['Có thể dùng ngay', 'Chưa nên xem là sự thật tuyệt đối'], [
        ('Các lỗi audit lặp lại 5/5 như heading-order, image-alt, link-name, page-error.', 'Finding do mô hình tự viết chỉ xuất hiện một lần.'),
        ('Tràn ngang được ghi nhận 4/5 lần.', 'Điểm “đẹp”, “dễ dùng” hoặc “chuyển đổi tốt” khi chưa có người dùng thật đối chiếu.'),
        ('Điểm tiêu chí trung bình và xu hướng lặp lại giữa 5 lần.', 'Điểm technical cao được hiểu là toàn bộ website tốt.'),
        ('Các hướng cần kiểm tra thủ công: menu, khoảng trống mobile, CTA và lỗi runtime.', 'Kết luận rằng AI không ảo giác chỉ vì hallucination proxy bằng 0%.'),
    ], widths=[3.45, 3.95])

    doc.add_heading('7. Phân tích ưu tiên cải thiện website ATAD', level=1)
    add_table(doc, ['Mức ưu tiên', 'Việc nên làm', 'Lý do'], [
        ('1', 'Điều tra lỗi page-error trên desktop và mobile bằng log trình duyệt.', 'Lỗi runtime lặp lại 5/5 có thể ảnh hưởng trực tiếp đến người dùng.'),
        ('2', 'Sửa các lỗi khả năng tiếp cận lặp lại: cấu trúc tiêu đề, alt ảnh, tên liên kết, landmark và tương phản.', 'Đây là nhóm có bằng chứng tự động và tái diễn rõ rệt.'),
        ('3', 'Kiểm tra tràn ngang desktop.', 'Đã xuất hiện 4/5 lần, có khả năng là lỗi bố cục thực.'),
        ('4', 'Rút gọn hoặc tổ chức lại menu lớn, làm rõ nội dung và hành động chính trong vùng nhìn đầu tiên.', 'Điểm visual thấp và nhiều finding mô tả cùng một hiện tượng.'),
        ('5', 'Rà soát mobile: khoảng trống đầu trang, mật độ chữ và khả năng đọc.', 'Responsive và readability là hai khu vực dao động hoặc bị chấm thấp.'),
        ('6', 'Làm nổi bật hành động liên hệ hoặc yêu cầu tư vấn.', 'Điểm conversion ổn định ở mức trung bình, chưa phải tín hiệu tốt.'),
    ], widths=[0.75, 3.55, 3.1])

    doc.add_heading('8. Phân tích ưu tiên cải thiện bộ máy đánh giá', level=1)
    add_manual_number(doc, 1, 'Chuẩn hóa ruleId và taxonomy cho finding do mô hình tạo ra, thay vì cho phép mô hình tự đặt tên hoàn toàn tự do.')
    add_manual_number(doc, 2, 'Bổ sung semantic matching để nhận ra các câu khác nhau nhưng cùng mô tả một hiện tượng.')
    add_manual_number(doc, 3, 'Tách recurrence của lỗi tự động, có thể kiểm chứng, khỏi recurrence của finding chủ quan do mô hình viết.')
    add_manual_number(doc, 4, 'Siết rubric cho các tiêu chí dao động như readability và visual.consistency bằng anchor rõ hơn và ví dụ điểm 2 so với điểm 3.')
    add_manual_number(doc, 5, 'Tạo bộ nhãn chuẩn do ít nhất hai reviewer đánh giá độc lập để đo calibration, tức mức độ máy khớp với con người.')
    add_manual_number(doc, 6, 'Khi chạy live, tiếp tục dùng retry theo từng stage, timeout dài và chế độ direct/offline khi gateway không ổn định.')

    doc.add_heading('9. Giải thích các thuật ngữ chính', level=1)
    add_table(doc, ['Thuật ngữ', 'Giải thích bằng ngôn ngữ thường'], [
        ('Rubric', 'Bộ tiêu chí và quy tắc dùng để chấm website.'),
        ('Criterion', 'Một tiêu chí nhỏ trong rubric, ví dụ khả năng đọc hoặc khoảng cách.'),
        ('Dimension', 'Một nhóm lớn gom nhiều criterion, ví dụ Visual hoặc UX.'),
        ('Finding', 'Một phát hiện cụ thể, ví dụ thiếu alt cho ảnh hoặc tràn ngang.'),
        ('Evidence', 'Bằng chứng đi kèm như ảnh chụp, kết quả audit hoặc log.'),
        ('Grounding', 'Mức độ một nhận xét có bám vào bằng chứng thật hay không.'),
        ('Hallucination proxy', 'Chỉ báo gián tiếp về việc mô hình có nêu điều không được bằng chứng hỗ trợ.'),
        ('Standard deviation', 'Độ dao động của điểm giữa các lần chạy; càng thấp càng ổn định.'),
        ('Exact agreement', 'Tỷ lệ các lần cho đúng cùng một điểm.'),
        ('Within-one agreement', 'Tỷ lệ các lần chênh không quá một bậc điểm.'),
        ('Recurrence', 'Tỷ lệ các lần một finding được nhận diện lại.'),
        ('Calibration', 'Đo mức độ điểm máy khớp với nhãn do con người chấm.'),
        ('Trust gate', 'Cửa kiểm soát tổng hợp; chỉ đạt khi mọi điều kiện tin cậy quan trọng đều đạt.'),
    ], widths=[1.7, 5.7])

    doc.add_heading('10. Kết luận cuối cùng', level=1)
    doc.add_paragraph('Kết quả 5 lần đánh giá cho thấy ATAD có các tín hiệu cần ưu tiên xử lý: lỗi runtime lặp lại, các lỗi khả năng tiếp cận lặp lại, nguy cơ tràn ngang, bố cục desktop bị menu lấn át, khoảng trống mobile lớn và hành động liên hệ chưa đủ nổi bật. Điểm kỹ thuật cao không phủ nhận các vấn đề về giao diện và trải nghiệm.')
    doc.add_paragraph('Ở góc độ bộ máy đánh giá, kết quả là khá tốt về tính ổn định của điểm số, độ đầy đủ của rubric, tính hợp lệ của số audit và việc gắn bằng chứng. Điểm chưa đạt nằm ở cách mô hình đặt tên và gom các finding bằng ngôn ngữ tự do. Vì vậy, kết luận phù hợp nhất là: có thể dùng báo cáo này để lập danh sách kiểm tra và ưu tiên cải thiện, nhưng cần người xác minh các nhận xét chủ quan trước khi dùng làm quyết định tự động.')

    doc.add_page_break()
    doc.add_heading('Phụ lục Nguồn dữ liệu', level=1)
    doc.add_paragraph('Các số liệu trong tài liệu được diễn giải từ bộ kết quả offline của ATAD với phiên bản chỉ số meta-eval-v2, mô hình coding-v3, prompt-v1 và rubric-v1. Hồ sơ kỹ thuật đầy đủ nằm trong các tệp sau:')
    add_bullet(doc, 'evals/promptfoo/results/atad-offline/report.md')
    add_bullet(doc, 'evals/promptfoo/results/atad-offline/meta-eval.json')
    add_bullet(doc, 'evals/promptfoo/results/atad-offline/run-ids.json')
    add_bullet(doc, 'docs/phan-tich-chi-so-meta-eval-atad.md')
    doc.add_paragraph('Thời điểm tạo bộ số liệu nguồn: 15/09/2026. Chế độ offline có nghĩa là bộ evidence đã thu thập được giữ lại để chấm lại; đây không phải là một lần truy cập trực tiếp mới vào website tại thời điểm mở tài liệu này.')

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run('WebLens | Phân tích kết quả đánh giá ATAD')
    set_run_font(fr, size=8.5, color=(120, 120, 120))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == '__main__':
    main()
